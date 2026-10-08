# -*- coding: utf-8 -*-
"""受領TXTの確認画面。生成画面の編集値は取り込まない。"""

import os
import csv
import io
import queue
import subprocess
import sys
import threading
from pathlib import Path

import tkinter as tk
from tkinter import filedialog, messagebox, ttk
from openpyxl import load_workbook

from .layout_inspect import (DETAIL_COLUMN_ORDER, DETAIL_DEFAULT_OPTIONAL_COLUMNS, DETAIL_HEADERS,
                             DETAIL_OPTIONAL_COLUMNS, DETAIL_REQUIRED_COLUMNS,
                             SCOPE, export_inspection, inspect_txt)


DETAIL_PREVIEW_WIDTHS = {0: 200, 1: 70, 2: 90, 5: 80, 6: 90, 7: 170,
                         8: 240, 9: 70, 10: 80, 11: 80, 12: 140,
                         14: 300, 20: 140, 21: 140, 22: 160, 23: 180}


def selected_detail_columns(optional_columns):
    selected = set(DETAIL_REQUIRED_COLUMNS) | set(optional_columns)
    return tuple(index for index in DETAIL_COLUMN_ORDER if index in selected)


class InspectionWindow(object):
    def __init__(self, parent=None, definition_options=None, encoding="cp932", output_dir=""):
        self.options = dict(definition_options or {})
        self.excel = self.options.pop("excel_path", "")
        self.result = None
        self.saved_path = None
        self.output_dir = output_dir
        self.events = queue.Queue()
        self.busy = False
        self.optional_detail_columns = set(DETAIL_DEFAULT_OPTIONAL_COLUMNS)
        self._cell_text = None
        self.window = tk.Toplevel(parent) if parent is not None else tk.Tk()
        self.window.title("Received TXT Checker — 受領TXT確認")
        width = min(1320, max(600, self.window.winfo_screenwidth() - 60))
        height = min(820, max(450, self.window.winfo_screenheight() - 100))
        self.window.geometry("%dx%d" % (width, height))
        self.window.columnconfigure(0, weight=1)
        self.window.rowconfigure(3, weight=1)
        self.encoding = tk.StringVar(value=encoding)
        self.status = tk.StringVar(value="受領TXTを選択してください。複数ファイルに対応します。")
        self.query = tk.StringVar(value="")
        self.issues_only = tk.BooleanVar(value=False)
        self.excel_label = tk.StringVar(value=str(self.excel) or "定義Excelを選択してください。")
        self.sheet = tk.StringVar(value=self.options.get("sheet_name") or "")
        source = ttk.LabelFrame(self.window, text="1. レイアウト定義Excel", padding=8)
        source.grid(row=0, column=0, sticky="we", padx=10, pady=8)
        source.columnconfigure(1, weight=1)
        self.excel_button = ttk.Button(source, text="Excelを選択...", command=self._choose_excel)
        self.excel_button.grid(row=0, column=0, padx=(0, 8))
        ttk.Label(source, textvariable=self.excel_label, wraplength=width - 220).grid(
            row=0, column=1, columnspan=3, sticky="w")
        ttk.Label(source, text="シート:").grid(row=1, column=0, pady=(8, 0))
        self.sheet_box = ttk.Combobox(source, textvariable=self.sheet, state="readonly", width=30)
        self.sheet_box.grid(row=1, column=1, sticky="w", pady=(8, 0))
        self.sheet_box.bind("<<ComboboxSelected>>", lambda _event: self._invalidate())
        self.columns_button = ttk.Button(source, text="見出し行・列設定...", command=self._column_settings)
        self.columns_button.grid(row=1, column=2, padx=8, pady=(8, 0))
        controls = ttk.Frame(self.window, padding=(10, 0))
        controls.grid(row=1, column=0, sticky="we")
        ttk.Label(controls, text="受領文字コード:").pack(side="left")
        self.encoding_box = ttk.Combobox(controls, textvariable=self.encoding,
                                         values=("cp932", "utf-8-sig", "utf-8"),
                                         state="readonly", width=12)
        self.encoding_box.pack(side="left", padx=6)
        self.read_button = ttk.Button(controls, text="受領TXTを選択・解析", command=self._choose)
        self.read_button.pack(side="left", padx=6)
        self.export_button = ttk.Button(controls, text="全件をExcel出力", command=self._export,
                                        state="disabled")
        self.export_button.pack(side="left", padx=6)
        self.export_columns_button = ttk.Button(controls, text="明細表示・出力項目...",
                                                 command=self._export_column_settings)
        self.export_columns_button.pack(side="left", padx=6)
        self.open_button = ttk.Button(controls, text="出力Excelを開く", command=self._open,
                                      state="disabled")
        self.open_button.pack(side="left", padx=6)
        filters = ttk.Frame(self.window, padding=10)
        filters.grid(row=2, column=0, sticky="we")
        ttk.Label(filters, text="検索（ファイル・FORM・項目名・実値）:").pack(side="left")
        entry = ttk.Entry(filters, textvariable=self.query, width=28)
        entry.pack(side="left", padx=6)
        entry.bind("<Return>", lambda _event: self._render())
        ttk.Checkbutton(filters, text="注意点のあるデータのみ", variable=self.issues_only,
                        command=self._render).pack(side="left", padx=6)
        ttk.Button(filters, text="検索", command=self._render).pack(side="left")
        notebook = ttk.Notebook(self.window)
        notebook.grid(row=3, column=0, sticky="nsew", padx=10)
        self.record_tree = self._tree(notebook, "レコード一覧", [
            "ファイル", "レコード", "FORM_ID", "対象有無", "参考情報", "備考"],
            [240, 75, 85, 75, 85, 580])
        indexes = selected_detail_columns(self.optional_detail_columns)
        self.detail_indexes = indexes
        self.detail_tree = self._tree(notebook, "項目明細", [DETAIL_HEADERS[i] for i in indexes],
                                     [DETAIL_PREVIEW_WIDTHS[i] for i in indexes])
        original = ttk.Frame(notebook)
        notebook.add(original, text="選択レコードの原文")
        original.rowconfigure(0, weight=1)
        original.columnconfigure(0, weight=1)
        self.raw_text = tk.Text(original, wrap="none", state="disabled")
        self.raw_text.grid(row=0, column=0, sticky="nsew")
        for orient, axis, row, col in (("vertical", "y", 0, 1), ("horizontal", "x", 1, 0)):
            bar = ttk.Scrollbar(original, orient=orient, command=getattr(self.raw_text, axis + "view"))
            bar.grid(row=row, column=col, sticky="ns" if axis == "y" else "ew")
            self.raw_text.configure(**{axis + "scrollcommand": bar.set})
        self.record_tree.bind("<<TreeviewSelect>>", self._show_raw)
        self.detail_tree.bind("<<TreeviewSelect>>", self._show_raw)
        ttk.Label(self.window, textvariable=self.status, wraplength=width - 30).grid(
            row=4, column=0, sticky="we", padx=10, pady=8)
        ttk.Label(self.window, text=SCOPE, wraplength=width - 30).grid(
            row=5, column=0, sticky="we", padx=10, pady=(0, 10))
        self.window.bind("<Button-1>", self._dismiss_cell, add="+")
        self.window.after(100, self._poll)
        if self.excel:
            self._load_sheets()

    def _invalidate(self):
        self.result = None
        self.saved_path = None
        self._clear()
        self.export_button.configure(state="disabled")
        self.open_button.configure(state="disabled")
        self.status.set("定義設定が変わりました。受領TXTを選択して再解析してください。")

    def _choose_excel(self):
        if self.busy:
            return
        filename = filedialog.askopenfilename(parent=self.window, title="レイアウト定義Excelを選択",
                                              filetypes=[("Excel", "*.xlsx *.xlsm")])
        if filename:
            self.excel = Path(filename)
            self.excel_label.set(str(self.excel))
            self.sheet.set("")
            self._invalidate()
            self._load_sheets()

    def _load_sheets(self):
        self.sheet_box.configure(values=())
        try:
            wb = load_workbook(str(self.excel), read_only=True, data_only=True)
            try:
                names = list(wb.sheetnames)
                active = wb.active.title
            finally:
                wb.close()
            self.sheet_box.configure(values=names)
            if self.sheet.get() not in names:
                self.sheet.set(active)
        except Exception as exc:
            self.sheet.set("")
            messagebox.showerror("定義Excelを読めません", str(exc), parent=self.window)

    def _sync_definition(self):
        if not self.excel or not Path(self.excel).is_file() or not self.sheet.get():
            messagebox.showwarning("定義Excel未選択", "定義Excelとシートを選択してください。", parent=self.window)
            return False
        self.options["sheet_name"] = self.sheet.get()
        return True

    def _column_settings(self):
        if self.busy:
            return
        dialog = tk.Toplevel(self.window)
        dialog.title("定義Excelの読込設定")
        dialog.transient(self.window)
        dialog.grab_set()
        ttk.Label(dialog, text="列は auto・列記号・見出し名で指定します。", padding=10).grid(
            row=0, column=0, columnspan=2)
        roles = [("header_row", "見出し行（空欄=自動）", ""),
                 ("form_column", "FORM_ID", "auto"), ("layout_column", "LAYOUT_ID", "auto"),
                 ("field_column", "ELEMENT_ID", "auto"), ("item_column", "項目名", "auto"),
                 ("data_type_column", "データ型", "I"), ("ime_column", "IME", "J"),
                 ("max_digits_column", "最大桁数", "K"),
                 ("input_attribute_column", "入力属性", "auto"),
                 ("input_rule_column", "入力規則", "auto"), ("notes_column", "補足", "auto"),
                 ("output_example_column", "出力例", "auto")]
        variables = {}
        for index, (key, label, default) in enumerate(roles, 1):
            ttk.Label(dialog, text=label).grid(row=index, column=0, sticky="w", padx=10, pady=3)
            variables[key] = tk.StringVar(value=str(self.options.get(key) or default))
            ttk.Entry(dialog, textvariable=variables[key], width=26).grid(
                row=index, column=1, padx=10, pady=3)

        def save():
            values = {key: variable.get().strip() for key, variable in variables.items()}
            try:
                header = int(values["header_row"]) if values["header_row"] else None
                if header is not None and header < 1:
                    raise ValueError()
            except ValueError:
                messagebox.showerror("見出し行", "1以上の整数または空欄にしてください。", parent=dialog)
                return
            values["header_row"] = header
            self.options.update(values)
            self._invalidate()
            dialog.destroy()
        ttk.Button(dialog, text="適用", command=save).grid(row=len(roles) + 1, column=1, pady=10)

    def _export_column_settings(self):
        if self.busy:
            return
        dialog = tk.Toplevel(self.window)
        dialog.title("項目明細の表示・出力列")
        dialog.transient(self.window)
        dialog.grab_set()
        ttk.Label(dialog, text="必須: FieldID・受領OCR値・属性・座標", padding=10).grid(
            row=0, column=0, columnspan=2, sticky="w")
        ttk.Label(dialog, text="明細項目", padding=(10, 0)).grid(row=1, column=0, sticky="w")
        ttk.Label(dialog, text="定義Excel", padding=(10, 0)).grid(row=1, column=1, sticky="w")
        variables = {}
        for column, group in enumerate((DETAIL_OPTIONAL_COLUMNS[:8], DETAIL_OPTIONAL_COLUMNS[8:])):
            for row, index in enumerate(group, 2):
                variable = tk.BooleanVar(value=index in self.optional_detail_columns)
                variables[index] = variable
                ttk.Checkbutton(dialog, text=DETAIL_HEADERS[index], variable=variable).grid(
                    row=row, column=column, sticky="w", padx=14, pady=3)

        def save():
            self.optional_detail_columns = {index for index, variable in variables.items()
                                            if variable.get()}
            dialog.destroy()
            self._refresh_detail_columns()

        ttk.Button(dialog, text="適用", command=save).grid(
            row=10, column=1, sticky="e", padx=10, pady=10)

    def _refresh_detail_columns(self):
        indexes = selected_detail_columns(self.optional_detail_columns)
        self.detail_indexes = indexes
        columns = [str(index) for index in range(len(indexes))]
        self.detail_tree.configure(columns=columns)
        for column, index in zip(columns, indexes):
            self.detail_tree.heading(column, text=DETAIL_HEADERS[index])
            self.detail_tree.column(column, width=DETAIL_PREVIEW_WIDTHS[index], stretch=False)
        self._render()

    def _tree(self, notebook, title, headers, widths):
        frame = ttk.Frame(notebook)
        notebook.add(frame, text=title)
        frame.rowconfigure(0, weight=1)
        frame.columnconfigure(0, weight=1)
        columns = [str(i) for i in range(len(headers))]
        tree = ttk.Treeview(frame, columns=columns, show="headings", selectmode="extended")
        tree.grid(row=0, column=0, sticky="nsew")
        for col, title, width in zip(columns, headers, widths):
            tree.heading(col, text=title)
            tree.column(col, width=width, stretch=False)
        for orient, axis, row, col in (("vertical", "y", 0, 1), ("horizontal", "x", 1, 0)):
            bar = ttk.Scrollbar(frame, orient=orient,
                                command=lambda *args, axis=axis: self._scroll_table(tree, axis, *args))
            bar.grid(row=row, column=col, sticky="ns" if axis == "y" else "ew")
            tree.configure(**{axis + "scrollcommand": bar.set})
        tree.tag_configure("注意点あり", background="#fff2cc")
        tree.bind("<Control-c>", lambda _event: self._copy_rows(tree))
        tree.bind("<Control-a>", lambda _event: self._select_all_rows(tree))
        if sys.platform == "darwin":
            tree.bind("<Command-c>", lambda _event: self._copy_rows(tree))
            tree.bind("<Command-a>", lambda _event: self._select_all_rows(tree))
        tree.bind("<Button-1>", self._show_cell)
        tree.bind("<B1-Motion>", self._drag_cell)
        tree.bind("<ButtonRelease-1>", self._release_cell)
        for sequence in ("<Configure>", "<MouseWheel>", "<Button-4>", "<Button-5>"):
            tree.bind(sequence, lambda _event: self._close_cell())
        return tree

    def _select_all_rows(self, tree):
        tree.selection_set(tree.get_children())
        return "break"

    def _copy_rows(self, tree):
        selected = set(tree.selection())
        if not selected:
            self.status.set("コピーする行を選択してください。セル内の文字をドラッグして選択できます。")
            return "break"
        output = io.StringIO(newline="")
        writer = csv.writer(output, delimiter="\t", lineterminator="\r\n")
        for iid in tree.get_children():
            if iid in selected:
                writer.writerow(tree.item(iid, "values"))
        self.window.clipboard_clear()
        self.window.clipboard_append(output.getvalue())
        self.status.set("%d行をコピーしました。Excelなどへ貼り付けできます。" % len(selected))
        return "break"

    def _close_cell(self):
        text = self._cell_text
        self._cell_text = None
        if text is not None:
            text.destroy()

    def _dismiss_cell(self, event):
        text = self._cell_text
        if text is not None and event.widget is not text:
            self._close_cell()

    def _scroll_table(self, tree, axis, *args):
        self._close_cell()
        getattr(tree, axis + "view")(*args)

    def _show_cell(self, event):
        tree = event.widget
        self._close_cell()
        # Shift/Ctrl（macOSはCommand）による複数行選択はTreeviewに任せる。
        modifiers = 0x0005 | (0x0008 if sys.platform == "darwin" else 0)
        if event.state & modifiers:
            return
        iid = tree.identify_row(event.y)
        column = tree.identify_column(event.x)
        if tree.identify_region(event.x, event.y) != "cell" or not iid or column == "#0":
            return
        tree.selection_set(iid)
        tree.focus(iid)
        text = self._select_cell_text(tree, iid, column)
        if text is not None:
            x, y, _width, _height = tree.bbox(iid, column)
            text.event_generate("<Button-1>", x=event.x - x, y=event.y - y)
        return "break"

    def _select_cell_text(self, tree, iid, column):
        self._close_cell()
        bounds = tree.bbox(iid, column)
        if not bounds:
            return None
        x, y, width, height = bounds
        background = "#fff2cc" if "注意点あり" in tree.item(iid, "tags") else "white"
        text = tk.Text(tree, wrap="none", exportselection=False, borderwidth=0,
                       highlightthickness=1, highlightbackground="#849bb7",
                       background=background, padx=2, pady=0,
                       font=ttk.Style(tree).lookup("Treeview", "font") or "TkDefaultFont")
        text.insert("1.0", tree.set(iid, column))
        text.configure(state="disabled")
        text.place(x=x, y=y, width=width, height=height)
        self._cell_text = text

        def select_all(_event):
            text.tag_add("sel", "1.0", "end-1c")
            return "break"

        def close(_event):
            self._close_cell()
            tree.focus_set()
            return "break"

        text.bind("<Control-a>", select_all)
        if sys.platform == "darwin":
            text.bind("<Command-a>", select_all)
        text.bind("<Escape>", close)
        text.focus_set()
        return text

    def _drag_cell(self, event):
        text = self._cell_text
        if text is not None and text.master is event.widget:
            text.event_generate("<B1-Motion>",
                                x=event.x - text.winfo_x(), y=event.y - text.winfo_y())
            return "break"

    def _release_cell(self, event):
        text = self._cell_text
        if text is not None and text.master is event.widget:
            text.event_generate("<ButtonRelease-1>",
                                x=event.x - text.winfo_x(), y=event.y - text.winfo_y())
            return "break"

    def _start(self, operation, callback):
        self.busy = True
        for widget in (self.read_button, self.export_button, self.open_button, self.encoding_box,
                       self.excel_button, self.sheet_box, self.columns_button,
                       self.export_columns_button):
            widget.configure(state="disabled")
        self.status.set("解析中..." if operation == "read" else "Excel保存中...")

        def worker():
            try:
                self.events.put((operation, callback(), None))
            except Exception as exc:
                self.events.put((operation, None, str(exc)))
        threading.Thread(target=worker, daemon=True).start()

    def _choose(self):
        if self.busy:
            return
        if not self._sync_definition():
            return
        paths = filedialog.askopenfilenames(parent=self.window, title="受領TXTを選択",
                                            filetypes=[("TXT", "*.txt"), ("全ファイル", "*")])
        if not paths:
            return
        self.result = None
        self.saved_path = None
        self._clear()
        encoding = self.encoding.get()
        self._start("read", lambda: inspect_txt(self.excel, paths, encoding=encoding, **self.options))

    def _export(self):
        if self.busy or self.result is None:
            return
        filename = filedialog.asksaveasfilename(
            parent=self.window, title="確認用Excelを保存", defaultextension=".xlsx",
            initialdir=self.output_dir or str(Path(self.excel).parent),
            initialfile="受領TXT確認.xlsx", filetypes=[("Excel", "*.xlsx")])
        if filename:
            # 同名ファイルの確認はネイティブ保存ダイアログが担当する。
            self.saved_path = None
            columns = selected_detail_columns(self.optional_detail_columns)
            self._start("export", lambda: export_inspection(
                self.result, filename, overwrite=True, detail_columns=columns))

    def _poll(self):
        try:
            operation, value, error = self.events.get_nowait()
        except queue.Empty:
            pass
        else:
            self.busy = False
            self.read_button.configure(state="normal")
            self.encoding_box.configure(state="readonly")
            self.excel_button.configure(state="normal")
            self.sheet_box.configure(state="readonly")
            self.columns_button.configure(state="normal")
            self.export_columns_button.configure(state="normal")
            if error:
                self.status.set("処理失敗: %s" % error)
                messagebox.showerror("受領TXT確認", error, parent=self.window)
            elif operation == "read":
                self.result = value
                self._render()
            else:
                self.saved_path = value
                self.status.set("全件保存完了: %s" % value)
            self.export_button.configure(state="normal" if self.result else "disabled")
            self.open_button.configure(state="normal" if self.saved_path else "disabled")
        self.window.after(100, self._poll)

    def _clear(self):
        self._close_cell()
        for tree in (self.record_tree, self.detail_tree):
            for iid in tree.get_children():
                tree.delete(iid)
        self.raw_text.configure(state="normal")
        self.raw_text.delete("1.0", "end")
        self.raw_text.configure(state="disabled")

    def _render(self):
        if self.busy or self.result is None:
            return
        self._clear()
        query = self.query.get().casefold()
        only = self.issues_only.get()
        visible = set()
        matched = 0
        shown = 0
        for index, row in enumerate(self.result.details):
            if only and not row[14]:
                continue
            if query and query not in " ".join(str(v) for v in row[:26]).casefold():
                continue
            matched += 1
            visible.add((row[0], row[1]))
            if shown < 3000:
                self.detail_tree.insert("", "end", iid="d%d" % index,
                                        values=["" if row[i] is None else row[i]
                                                for i in self.detail_indexes], tags=(row[13],))
                shown += 1
        record_shown = 0
        record_matched = 0
        for index, record in enumerate(self.result.records):
            if only and not record["issues"]:
                continue
            if query and (record["file"], record["number"]) not in visible and query not in (
                    record["file"] + " " + record["form"] + " " + record["issues"]).casefold():
                continue
            record_matched += 1
            if record_shown < 3000:
                self.record_tree.insert("", "end", iid="r%d" % index,
                                        values=[record[k] for k in (
                                            "file", "number", "form", "presence", "status", "issues")],
                                        tags=(record["status"],))
                record_shown += 1
        count = sum(bool(r["issues"]) for r in self.result.records)
        self.status.set("全%dレコード / 注意点のあるレコード%d。表示: レコード%d/%d・明細%d/%d。"
                        "画面は各3000件まで。Excel出力は検索条件に関係なく全件。文字コード: %s" % (
                            len(self.result.records), count,
                            record_shown, record_matched, shown, matched, self.result.encoding))

    def _show_raw(self, event):
        selected = event.widget.selection()
        if not selected or self.result is None:
            return
        iid = selected[0]
        if iid.startswith("r"):
            record = self.result.records[int(iid[1:])]
        else:
            row = self.result.details[int(iid[1:])]
            record = next(r for r in self.result.records if (r["file"], r["number"]) == (row[0], row[1]))
        self.raw_text.configure(state="normal")
        self.raw_text.delete("1.0", "end")
        self.raw_text.insert("1.0", record["raw"])
        self.raw_text.configure(state="disabled")

    def _open(self):
        if not self.saved_path:
            return
        try:
            if sys.platform == "win32":
                os.startfile(str(self.saved_path))
            else:
                subprocess.Popen(["open" if sys.platform == "darwin" else "xdg-open", str(self.saved_path)])
        except OSError as exc:
            messagebox.showerror("Excelを開けません", str(exc), parent=self.window)


def main(initial_excel=None):
    options = {"excel_path": initial_excel} if initial_excel else None
    app = InspectionWindow(definition_options=options)
    app.window.mainloop()
    return 0


if __name__ == "__main__":
    sys.exit(main())
