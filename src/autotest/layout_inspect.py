# -*- coding: utf-8 -*-
"""受領 TXT を定義と照合し、原値を保持した確認用 Excel を作る。"""

import argparse
import csv
import hashlib
import io
import os
import re
import sys
import tempfile
import zipfile
from collections import OrderedDict
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .layout_txt import LayoutTxtError, read_layout_fields


SCOPE = ("Excel定義に対応する実値を表示し、項目数・文字数などの注意点を参考情報として記載します。"
         "座標は任意です。空欄でも注意点にしません。")
DETAIL_HEADERS = [
    "ファイル", "レコード", "物理開始行", "FORM_ID", "対象有無", "TXT項目順",
    "FieldID", "項目名", "受領OCR値", "文字数", "最大桁数", "属性", "座標",
    "参考情報", "備考", "定義シート", "定義行", "LAYOUT_ID", "ELEMENT_ID",
    "展開番号", "データ型", "IME", "入力属性", "入力規則", "補足", "出力例",
    "確認内容（手入力）", "確認者", "確認メモ",
]
# 定義との対応付けは位置で行う。FieldIDは受領値としてのみ出力する。
DETAIL_REQUIRED_COLUMNS = (6, 8, 11, 12)
DETAIL_DEFAULT_OPTIONAL_COLUMNS = (0, 1, 2, 5, 7, 14)
DETAIL_OPTIONAL_COLUMNS = (0, 1, 2, 5, 7, 9, 10, 14, 20, 21, 22, 23, 24, 25)
# 備考は選択時も末尾に置く。
DETAIL_COLUMN_ORDER = (0, 1, 2, 5, 6, 7, 8, 9, 10, 11, 12,
                       20, 21, 22, 23, 24, 25, 14)
DETAIL_OUTPUT_COLUMNS = tuple(index for index in DETAIL_COLUMN_ORDER
                              if index in DETAIL_REQUIRED_COLUMNS + DETAIL_DEFAULT_OPTIONAL_COLUMNS)
DETAIL_COLUMN_WIDTHS = {0: 42, 1: 12, 2: 12, 5: 12, 6: 14, 7: 26,
                        8: 42, 9: 10, 10: 12, 11: 12, 12: 23, 14: 65,
                        20: 20, 21: 20, 22: 24, 23: 30, 24: 42, 25: 42}
RECEIVED_ATTRIBUTE_VALUES = ("0", "4", "8", "12")


def _sha(data):
    return hashlib.sha256(data).hexdigest()


class Inspection(object):
    def __init__(self, definition, sheet, header, columns, definition_hash, encoding):
        self.definition = str(definition)
        self.sheet = sheet
        self.header = header
        self.columns = columns
        self.definition_hash = definition_hash
        self.encoding = encoding
        self.created = datetime.now().isoformat(timespec="seconds")
        self.files = []
        self.records = []
        self.details = []

    @property
    def exit_code(self):
        # 参考情報の有無は、読込・Excel出力の成功とは別に扱う。
        return 0


def _detail(result, record, position, block, field, issues, status):
    field_id, value, attribute, coordinates = (block + [None] * 4)[:4]
    return [
        record["file"], record["number"], record["line"], record["form"],
        record["presence"], position, field_id,
        field.item_name if field else "", value, len(value) if value is not None else None,
        field.max_digits if field else None, attribute, coordinates, status,
        " / ".join(issues), result.sheet if field else "", field.row_number if field else None,
        field.layout_id if field else "", field.source_element_id if field else "",
        ("%d/%d" % (field.occurrence_index, field.occurrence_count)) if field else "",
        field.data_type if field else "", field.ime_name if field else "",
        field.input_attribute if field else "", field.input_rule if field else "",
        field.notes if field else "", field.output_example if field else "", "", "", "",
    ]


def _detect_block_width(values, fields):
    if len(values) <= 2:
        return 4
    scores = {}
    for width in (3, 4):
        blocks = [values[index:index + width] for index in range(2, len(values), width)]
        scores[width] = (
            sum(len(block) > 2 and block[2] not in RECEIVED_ATTRIBUTE_VALUES for block in blocks),
            sum(len(block) < 3 for block in blocks),
            abs(len(blocks) - len(fields)) if fields else 0,
        )
    if scores[3] == scores[4]:
        # 1項目・座標なしでは両方式の読み取り内容が同じなので曖昧さはない。
        if len(values) <= 5:
            return 3
        # OCR値を別項目のIDなどに誤って割り当てるより原文を残す。
        return None
    return min(scores, key=scores.get)


def _inspect_record(result, groups, record, values, parse_error, block_width=None):
    issues = []
    review = []
    record.update(form=values[0] if values else "",
                  presence=values[1] if len(values) > 1 else "")
    fields = groups.get(record["form"])
    if parse_error:
        issues.append(parse_error)
    if len(values) < 2:
        issues.append("FORM_ID / 対象有無が不足")
    if fields is None:
        issues.append("FORM_IDが定義外または空")
    if record["presence"] not in ("0", "1"):
        issues.append("対象有無は0または1が必要")
    elif record["presence"] == "0":
        review.append("対象有無=0。項目の要否を確認")
    width = block_width or _detect_block_width(values, fields)
    uncertain = width is None
    if uncertain:
        issues.append("項目の区切りを特定できません（原文を表示）")
    blocks = [values[index:index + width] for index in range(2, len(values), width)] if width else []
    if fields and not uncertain and len(blocks) != len(fields):
        issues.append("項目数不一致: 定義%d / 受領%d" % (len(fields), len(blocks)))
    detail_statuses = []
    for index, block in enumerate(blocks, 1):
        field = fields[index - 1] if fields and index <= len(fields) else None
        problems = []
        cautions = []
        if fields and field is None:
            problems.append("定義に対応項目なし（受領の余剰項目）")
        if len(block) < 3:
            problems.append("項目ブロックのOCR値または属性が不足")
        value = block[1] if len(block) > 1 else None
        if value is not None and field and field.max_digits is not None:
            if len(value) > field.max_digits:
                problems.append("最大桁数超過（文字数基準）")
        if value == "":
            cautions.append("OCR値が空")
        if len(block) > 2:
            if block[2] not in RECEIVED_ATTRIBUTE_VALUES:
                problems.append("属性が定義外")
        if len(block) > 3 and block[3] and not re.fullmatch(r"[0-9]+,[0-9]+,[0-9]+,[0-9]+", block[3]):
            cautions.append("座標の形式を確認（参考:4個の非負整数）")
        status = "注意点あり" if problems or cautions else ""
        detail_statuses.append(status)
        result.details.append(_detail(result, record, index, block, field,
                                      problems + cautions, status))
    for field in (fields[len(blocks):] if fields and not uncertain else []):
        result.details.append(_detail(result, record, None, [None], field,
                                      ["TXTに項目なし（定義から補記）"], "注意点あり"))
        detail_statuses.append("注意点あり")
    if "注意点あり" in detail_statuses:
        review.append("項目明細に注意点あり")
    record.update(status="注意点あり" if issues or review else "",
                  issues=" / ".join(issues + review), expected=len(fields) if fields else None,
                  actual=len(blocks) if not uncertain else None)
    result.records.append(record)


def inspect_txt(excel_path, txt_paths, encoding="cp932", block_width=None, **definition_options):
    """同一FORMの複数レコードも別々に照合し、実値を補正しない。"""
    definition = Path(excel_path).resolve()
    paths = [Path(path).resolve() for path in txt_paths]
    if not paths:
        raise LayoutTxtError("受領TXTを1件以上選択してください。")
    if len(set(paths)) != len(paths):
        raise LayoutTxtError("同じ受領TXTが重複選択されています。")
    if block_width not in (None, 3, 4):
        raise LayoutTxtError("項目形式は座標列あり（4要素）または座標列なし（3要素）にしてください。")
    # 生成用既定値・画面編集は実データの解釈に混ぜない。
    definition_options.update(default_value_column="none", coordinates_column="none",
                              profile="normal", date_mode="wareki", coverage_form_id="")
    try:
        before = definition.read_bytes()
        fields, sheet, header, columns = read_layout_fields(definition, **definition_options)
        if definition.read_bytes() != before:
            raise LayoutTxtError("読込中に定義Excelが変更されました。再実行してください。")
    except OSError as exc:
        raise LayoutTxtError("定義Excelを読めません: %s" % exc)
    result = Inspection(definition, sheet, header, columns, _sha(before), encoding)
    result.block_width = block_width
    groups = OrderedDict()
    for field in fields:
        groups.setdefault(field.form_id, []).append(field)
    for path in paths:
        try:
            payload = path.read_bytes()
            text = payload.decode("utf-8" if encoding == "utf-8-sig" else encoding)
        except (OSError, UnicodeError, LookupError) as exc:
            raise LayoutTxtError("受領TXTを読めません（文字コード %s）: %s: %s" %
                                 (encoding, path, exc))
        result.files.append((str(path), len(payload), _sha(payload)))
        # BOMは解析時だけ除去する。証跡の原文とハッシュには保持する。
        raw_lines = io.StringIO(text, newline="").readlines()
        parsed_text = text[1:] if text.startswith("\ufeff") else text
        reader = csv.reader(io.StringIO(parsed_text, newline=""), strict=True)
        previous = 0
        number = 0
        while True:
            error = ""
            try:
                values = next(reader)
            except StopIteration:
                break
            except csv.Error as exc:
                # 不正引用符で境界を失った後は推測して次レコードに割り当てない。
                values = []
                error = "CSV解析不可: %s（以降の原文を保持）" % exc
            number += 1
            end = len(raw_lines) if error else reader.line_num
            record = dict(file=str(path), number=number, line=previous + 1,
                          end_line=end, raw="".join(raw_lines[previous:end]))
            _inspect_record(result, groups, record, values, error, block_width=block_width)
            previous = end
            if error:
                break
        if not number:
            _inspect_record(result, groups, dict(file=str(path), number=1, line=1,
                                                end_line=1, raw=text), [], "空ファイル", block_width=block_width)
    return result


def _excel_text(value):
    # XMLに格納できない制御文字は可視化する。通常の値は変更しない。
    return re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ufffe\uffff]",
                  lambda match: "\\u%04X" % ord(match.group()), value)


def _append(ws, values):
    if ws.max_row >= 1048576:
        raise LayoutTxtError("Excelの行数上限を超えました。TXTを分割して実行してください。")
    row = ws.max_row + 1 if ws.max_row > 1 or ws.cell(1, 1).value is not None else 1
    for col, value in enumerate(values, 1):
        if isinstance(value, str):
            value = _excel_text(value)
            if len(value) > 32767:
                raise LayoutTxtError("Excelのセル文字数上限を超えました: %s 行%d 列%d" %
                                     (ws.title, row, col))
        cell = ws.cell(row, col, value)
        if isinstance(value, str):
            # =始まりや先頭ゼロも受領文字列そのもの。数式として実行しない。
            cell.data_type = "s"
            cell.number_format = "@"


def _style(ws, widths, status_column=None):
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.sheet_view.showGridLines = False
    line = Side(style="thin", color="B6C2CD")
    border = Border(left=line, right=line, top=line, bottom=line)
    for index, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(index)].width = width
    for cell in ws[1]:
        cell.border = border
        cell.fill = PatternFill("solid", fgColor="244660")
        cell.font = Font(name="Meiryo", size=10, color="FFFFFF", bold=True)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 32
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.border = border
            cell.font = Font(name="Meiryo", size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        ws.row_dimensions[row[0].row].height = 45
        if status_column:
            cell = row[status_column - 1]
            if cell.value == "注意点あり":
                cell.fill = PatternFill("solid", fgColor="FFF2CC")


def export_inspection(result, output_path, overwrite=False, detail_columns=None):
    """原本を保護し、保存失敗時にも既存成果物を壊さない。"""
    selected = set(DETAIL_OUTPUT_COLUMNS if detail_columns is None else detail_columns)
    allowed = set(DETAIL_COLUMN_ORDER)
    if not set(DETAIL_REQUIRED_COLUMNS) <= selected or not selected <= allowed:
        raise LayoutTxtError("項目明細の出力列が不正です。")
    columns = tuple(index for index in DETAIL_COLUMN_ORDER if index in selected)
    path = Path(output_path).resolve()
    sources = [Path(result.definition)] + [Path(item[0]) for item in result.files]
    if path in sources or (path.exists() and any(os.path.samefile(str(path), str(p))
                                               for p in sources if p.exists())):
        raise LayoutTxtError("出力先に入力原本は指定できません。")
    if path.suffix.lower() != ".xlsx":
        raise LayoutTxtError("出力先は .xlsx にしてください。")
    if path.exists() and not overwrite:
        raise LayoutTxtError("同名成果物があります。別名にするか上書きを指定してください。")
    wb = Workbook()
    details = wb.active
    details.title = "項目明細"
    _append(details, [DETAIL_HEADERS[index] for index in columns])
    for row in result.details:
        _append(details, [row[index] for index in columns])
    _style(details, [DETAIL_COLUMN_WIDTHS[index] for index in columns])
    details.freeze_panes = "B2"
    if 14 in columns:
        note_column = columns.index(14) + 1
        for row in details.iter_rows(min_row=2, min_col=note_column, max_col=note_column):
            if row[0].value:
                row[0].fill = PatternFill("solid", fgColor="FFF2CC")
    records = wb.create_sheet("レコード一覧")
    _append(records, ["ファイル", "レコード", "物理開始行", "物理終了行", "FORM_ID", "対象有無",
                      "定義項目数", "受領項目数", "参考情報", "備考"])
    for record in result.records:
        _append(records, [record[key] for key in (
            "file", "number", "line", "end_line", "form", "presence", "expected",
            "actual", "status", "issues")])
    _style(records, [42, 12, 12, 12, 14, 12, 14, 14, 14, 85], 9)
    raw = wb.create_sheet("受領原文")
    _append(raw, ["ファイル", "レコード", "物理開始行", "分割番号", "原文（分割順・制御文字は可視化）"])
    for record in result.records:
        chunks = [record["raw"][i:i + 10000] for i in range(0, len(record["raw"]), 10000)] or [""]
        for index, chunk in enumerate(chunks, 1):
            # 制御文字の可視化後もセル上限を守る。
            visible = _excel_text(chunk)
            if len(visible) > 32767:
                raise LayoutTxtError("原文の制御文字が多すぎるためExcelに格納できません。")
            _append(raw, [record["file"], record["number"], record["line"], index, visible])
    _style(raw, [42, 12, 12, 12, 115])
    path.parent.mkdir(parents=True, exist_ok=True)
    fd, temp_name = tempfile.mkstemp(prefix=".inspection_", suffix=".xlsx", dir=str(path.parent))
    os.close(fd)
    try:
        packed = io.BytesIO()
        # Windowsの旧XML writerによる改行変換を避ける。1文字の置換なので
        # セル上限近くの文字列でもopenpyxlによる切り詰めを起こさない。
        text_cells = [cell for sheet in wb for row in sheet for cell in row
                      if isinstance(cell.value, str)]
        used = set("".join(cell.value for cell in text_cells))
        markers = []
        for codepoint in range(0xF0000, 0xFFFFE):
            character = chr(codepoint)
            if character not in used:
                markers.append(character)
                if len(markers) == 2:
                    break
        if len(markers) != 2:
            raise LayoutTxtError("改行保持用の文字を確保できません。")
        originals = []
        try:
            for cell in text_cells:
                if "\r" in cell.value or "\n" in cell.value:
                    originals.append((cell, cell.value))
                    cell.value = cell.value.replace("\r", markers[0]).replace("\n", markers[1])
            wb.save(packed)
        finally:
            for cell, value in originals:
                cell.value = value
        packed.seek(0)
        # CR/LFを文字参照に戻し、XML読込時の改行正規化も防ぐ。
        with zipfile.ZipFile(packed, "r") as source_zip, zipfile.ZipFile(temp_name, "w") as target_zip:
            for entry in source_zip.infolist():
                data = source_zip.read(entry.filename)
                if entry.filename.endswith(".xml"):
                    for marker, reference in zip(markers, (b"&#13;", b"&#10;")):
                        data = data.replace(marker.encode("utf-8"), reference)
                        data = data.replace(("&#%d;" % ord(marker)).encode("ascii"), reference)
                target_zip.writestr(entry, data)
        if overwrite:
            os.replace(temp_name, str(path))
        else:
            # x モードにより保存直前の同名ファイル作成でも上書きしない。
            with open(temp_name, "rb") as src, path.open("xb") as dst:
                try:
                    import shutil
                    shutil.copyfileobj(src, dst)
                except Exception:
                    dst.close()
                    path.unlink()
                    raise
    finally:
        wb.close()
        if os.path.exists(temp_name):
            os.unlink(temp_name)
    return path


def main(argv=None):
    parser = argparse.ArgumentParser(description="受領TXTを定義Excelで照合して確認用Excelに出力")
    parser.add_argument("excel", help="レイアウト定義Excel")
    parser.add_argument("txt", nargs="+", help="受領TXT（複数指定可）")
    parser.add_argument("--out", required=True, help="確認用.xlsx出力先")
    parser.add_argument("--encoding", default="cp932", choices=("cp932", "utf-8", "utf-8-sig"))
    parser.add_argument("--sheet", default=None)
    parser.add_argument("--header-row", type=int, default=None)
    for role, default in (("form", "auto"), ("layout", "auto"), ("field", "auto"),
                          ("item", "auto"), ("data-type", "I"), ("ime", "J"),
                          ("max-digits", "K")):
        parser.add_argument("--%s-column" % role, default=default)
    parser.add_argument("--overwrite", action="store_true")
    parser.add_argument("--no-coordinates", action="store_true", help="座標列なしの3要素形式で読む")
    args = parser.parse_args(argv)
    options = {role + "_column": getattr(args, role + "_column") for role in
               ("form", "layout", "field", "item", "data_type", "ime", "max_digits")}
    try:
        result = inspect_txt(args.excel, args.txt, encoding=args.encoding,
                             block_width=3 if args.no_coordinates else None,
                             sheet_name=args.sheet, header_row=args.header_row, **options)
        path = export_inspection(result, args.out, overwrite=args.overwrite)
    except (LayoutTxtError, OSError, ValueError) as exc:
        print("[設定エラー] %s" % exc, file=sys.stderr)
        return 2
    print("確認用Excel: %s / レコード数: %d" % (path, len(result.records)))
    print(SCOPE)
    return result.exit_code


if __name__ == "__main__":
    sys.exit(main())
