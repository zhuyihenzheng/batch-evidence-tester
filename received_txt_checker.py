# -*- coding: utf-8 -*-
"""受領TXT確認専用アプリの起動・配布用エントリーポイント。"""

import argparse
import gc
import shutil
import sys
import tempfile
import time
from pathlib import Path

APP_VERSION = "0.1.15"


def _prepare_source_path():
    if not getattr(sys, "frozen", False):
        source = str(Path(__file__).resolve().parent / "src")
        if source not in sys.path:
            sys.path.insert(0, source)


def _cleanup_smoke_directory(path, attempts=10, delay=0.1):
    """Windowsの短時間の共有ロックだけ再試行し、実処理の結果を上書きしない。"""
    for attempt in range(attempts):
        gc.collect()
        try:
            shutil.rmtree(str(path))
            return True
        except OSError as exc:
            if not Path(path).exists():
                return True
            error_code = getattr(exc, "winerror", None) or exc.errno
            if error_code not in (32, 33):
                raise
            if attempt + 1 < attempts:
                time.sleep(delay)
    return False


def _functional_smoke_test():
    from openpyxl import Workbook, load_workbook
    from autotest.layout_inspect import inspect_txt, export_inspection

    root = Path(tempfile.mkdtemp(prefix="received_txt_checker_"))
    try:
        definition = root / "definition.xlsx"
        txt = root / "received.txt"
        output = root / "check.xlsx"
        wb = Workbook()
        ws = wb.active
        ws.append(["FORM_ID", "LAYOUT_ID", "", "", "", "", "", "ITEM_NAME",
                   "ELEMENT_DATA_TYPE_NAME", "ELEMENT_IME_NAME", "MAX_NUM_DIGITS", "ELEMENT_ID"])
        ws.append(["1001", "1", "", "", "", "", "", "受付番号", "文字列", "半角英数", 10, "9001"])
        wb.save(str(definition))
        wb.close()
        txt.write_bytes('"1001","1","1","000123","0","0,0,1,1"\r\n'.encode("cp932"))
        result = inspect_txt(definition, [txt])
        if result.exit_code != 0 or result.details[0][7] != "受付番号":
            raise RuntimeError("定義との照合に失敗しました。")
        export_inspection(result, output)
        # 古いopenpyxlでもOSのファイルハンドルを確実に閉じる。
        with output.open("rb") as stream:
            wb = load_workbook(stream, read_only=True, data_only=True)
            try:
                if (wb["項目明細"]["E2"].value != "1" or
                        wb["項目明細"]["G2"].value != "000123"):
                    raise RuntimeError("Excel出力で受領値が変わりました。")
            finally:
                wb.close()
    finally:
        processing_error = sys.exc_info()[0] is not None
        try:
            cleaned = _cleanup_smoke_directory(root)
        except OSError:
            if not processing_error:
                raise
            cleaned = False
        if not cleaned and sys.stderr is not None:
            print("一時ファイルを削除できませんでした: %s" % root, file=sys.stderr)
    return 0


def main(argv=None):
    parser = argparse.ArgumentParser(description="受領TXT確認専用アプリ")
    parser.add_argument("--version", action="version", version="ReceivedTxtChecker " + APP_VERSION)
    parser.add_argument("--smoke-test", action="store_true", help="画面を開かず読込・出力を検証")
    parser.add_argument("--excel", help="起動時に選択する定義Excel")
    args = parser.parse_args(argv)
    _prepare_source_path()
    if args.smoke_test:
        return _functional_smoke_test()
    from autotest.layout_inspect_gui import main as gui_main
    return gui_main(initial_excel=Path(args.excel) if args.excel else None)


if __name__ == "__main__":
    sys.exit(main())
