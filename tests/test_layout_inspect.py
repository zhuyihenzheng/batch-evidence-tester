# -*- coding: utf-8 -*-
"""受領TXTの照合で実値を失わず、異常を成功扱いしないこと。"""

import csv
import hashlib
import io
import shutil
import sys
import tempfile
import time
import unittest
import zipfile
from pathlib import Path
from unittest import mock

from openpyxl import Workbook, load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "src"))
from autotest.layout_inspect import export_inspection, inspect_txt, main
from autotest.layout_txt import LayoutTxtError, generate_layout_txt


class InspectionCase(unittest.TestCase):
    def setUp(self):
        self.root = Path(tempfile.mkdtemp(prefix="layout_inspection_test_"))
        self.definition = self.root / "definition.xlsx"
        self.txt = self.root / "received.txt"
        self.output = self.root / "inspection.xlsx"
        wb = Workbook()
        ws = wb.active
        ws.title = "定義"
        ws.append(["FORM_ID", "LAYOUT_ID", "", "", "", "", "", "ITEM_NAME",
                   "ELEMENT_DATA_TYPE_NAME", "ELEMENT_IME_NAME", "MAX_NUM_DIGITS", "ELEMENT_ID"])
        ws.append(["1001", "01", "", "", "", "", "", "受付番号", "文字列", "半角英数", 20, "0099"])
        ws.append(["1001", "01", "", "", "", "", "", "氏名", "文字列", "全タイプ", 20, "0100"])
        ws.append(["2001", "02", "", "", "", "", "", "処置日", "カレンダー", "", None, "0200"])
        wb.save(str(self.definition))
        wb.close()

    def tearDown(self):
        shutil.rmtree(str(self.root))

    def write_records(self, rows, encoding="cp932"):
        output = io.StringIO(newline="")
        writer = csv.writer(output, quoting=csv.QUOTE_ALL, lineterminator="\r\n")
        writer.writerows(rows)
        self.txt.write_bytes(output.getvalue().encode(encoding))

    def row(self, first="0000123", second="山田 太郎"):
        return ["1001", "1", "1", first, "0", "0,0,10,10", "2", second, "0", "0,0,10,10"]

    def inspect(self, **kwargs):
        return inspect_txt(self.definition, [self.txt], **kwargs)

    def test_actuals_follow_txt_position_even_when_field_ids_are_reversed(self):
        row = self.row()
        self.write_records([row[:2] + row[6:] + row[2:6]])
        result = self.inspect()
        self.assertEqual(result.exit_code, 0)
        self.assertEqual([r[7] for r in result.details], ["受付番号", "氏名"])
        self.assertEqual([r[18] for r in result.details], ["0099", "0100"])
        self.assertEqual([r[8] for r in result.details], ["山田 太郎", "0000123"])
        self.assertEqual(result.records[0]["issues"], "")

    def test_received_field_ids_starting_at_6001_use_one_set_of_detail_rows(self):
        row = self.row(first="3", second="手続き値")
        row[2], row[6] = "6001", "6002"
        self.write_records([row])
        result = self.inspect()
        self.assertEqual(result.records[0]["issues"], "")
        self.assertEqual(result.records[0]["actual"], 2)
        self.assertEqual(len(result.details), 2)
        self.assertEqual([detail[7] for detail in result.details], ["受付番号", "氏名"])
        self.assertEqual([detail[8] for detail in result.details], ["3", "手続き値"])
        self.assertEqual([detail[6] for detail in result.details], ["6001", "6002"])
        export_inspection(result, self.output)
        wb = load_workbook(str(self.output))
        try:
            self.assertEqual(wb["項目明細"].max_row, 3)
            self.assertEqual(wb["項目明細"]["E2"].value, "6001")
            self.assertEqual(wb["項目明細"]["G2"].value, "3")
        finally:
            wb.close()

    def test_one_received_field_starting_at_6002_fills_first_position(self):
        self.write_records([["1001", "1", "6002", "手続き値", "0", "0,0,0,0"]])
        result = self.inspect()
        self.assertEqual(len(result.details), 2)
        self.assertEqual(result.details[0][7:9], ["受付番号", "手続き値"])
        self.assertEqual(result.details[1][7], "氏名")
        self.assertIsNone(result.details[1][6])
        self.assertIsNone(result.details[1][8])
        self.assertIn("TXTに項目なし", result.details[1][14])

    def test_35_items_starting_at_6001_do_not_form_two_sections(self):
        wb = load_workbook(str(self.definition))
        try:
            for number in range(1, 36):
                wb.active.append(["5003", "01", "", "", "", "", "",
                                  "手続%d" % number, "文字列", "全タイプ", 20,
                                  str(8000 + number)])
            wb.save(str(self.definition))
        finally:
            wb.close()
        values = ["5003", "1"]
        for number in range(1, 36):
            values.extend([str(6000 + number), str(number), "0", "0,0,0,0"])
        self.write_records([values])
        result = self.inspect()
        self.assertEqual(result.records[0]["issues"], "")
        self.assertEqual(result.records[0]["actual"], 35)
        self.assertEqual(len(result.details), 35)
        self.assertEqual([row[7] for row in result.details],
                         ["手続%d" % number for number in range(1, 36)])
        export_inspection(result, self.output)
        exported = load_workbook(str(self.output))
        try:
            self.assertEqual(exported["項目明細"].max_row, 36)
            self.assertEqual(exported["項目明細"]["E36"].value, "6035")
            self.assertEqual(exported["項目明細"]["G36"].value, "35")
        finally:
            exported.close()

    def test_repeated_nonnumeric_field_ids_do_not_affect_mapping(self):
        row = self.row()
        row[2], row[6] = "X", "X"
        self.write_records([row])
        result = self.inspect()
        self.assertEqual(len(result.details), 2)
        self.assertEqual(result.records[0]["issues"], "")
        self.assertEqual([detail[7] for detail in result.details], ["受付番号", "氏名"])

    def test_calendar_expansion_can_use_6001_series(self):
        values = ["2001", "1"]
        for number in range(1, 47):
            values.extend([str(6000 + number), "5/8/6/1", "0", "0,0,0,0"])
        self.write_records([values])
        result = self.inspect()
        self.assertEqual(result.records[0]["issues"], "")
        self.assertEqual(result.records[0]["actual"], 46)
        self.assertEqual(len(result.details), 46)
        self.assertEqual(result.details[-1][19], "46/46")

    def test_mixed_thousand_series_still_uses_txt_position(self):
        row = self.row()
        row[2], row[6] = "6001", "7002"
        self.write_records([row])
        result = self.inspect()
        self.assertEqual(result.records[0]["issues"], "")
        self.assertEqual([detail[6] for detail in result.details[:2]], ["6001", "7002"])
        self.assertEqual([detail[7] for detail in result.details[:2]], ["受付番号", "氏名"])
        self.assertEqual(result.records[0]["raw"], self.txt.read_bytes().decode("cp932"))

    def test_multiline_quotes_spaces_leading_zero_and_formula_are_preserved(self):
        self.write_records([self.row('  "零",\r\n001 ', '=1+1'), self.row()])
        result = self.inspect()
        self.assertEqual(result.exit_code, 0)
        self.assertEqual(len(result.records), 2)
        self.assertEqual(result.records[1]["line"], 3)
        self.assertEqual(result.details[0][8], '  "零",\r\n001 ')
        export_inspection(result, self.output)
        wb = load_workbook(str(self.output))
        try:
            self.assertEqual(wb.sheetnames, ["項目明細", "レコード一覧", "受領原文"])
            headers = [cell.value for cell in wb["項目明細"][1]]
            self.assertEqual(headers, ["ファイル", "レコード", "物理開始行", "TXT項目順",
                                       "FieldID", "項目名", "受領OCR値", "属性", "座標", "備考"])
            for removed in ("FORM_ID", "対象有無", "最大桁数", "文字数", "定義シート", "確認者"):
                self.assertNotIn(removed, headers)
            self.assertEqual(wb["レコード一覧"]["E2"].value, "1001")
            self.assertEqual(wb["レコード一覧"]["F2"].value, "1")
            self.assertEqual(wb["項目明細"]["G2"].value, '  "零",\r\n001 ')
            self.assertEqual(wb["項目明細"]["G3"].value, '=1+1')
            self.assertEqual(wb["項目明細"]["G3"].data_type, "s")
            self.assertEqual(wb["項目明細"]["G4"].value, '0000123')
            self.assertEqual(wb["項目明細"].freeze_panes, "B2")
            self.assertEqual(wb["項目明細"].auto_filter.ref, "A1:J5")
            self.assertEqual(wb["受領原文"]["E2"].value + wb["受領原文"]["E3"].value,
                             self.txt.read_bytes().decode("cp932"))
        finally:
            wb.close()

    def test_optional_detail_columns_are_exported_only_when_selected(self):
        row = self.row(first="123456789012345678901")
        row[2] = "6001"
        self.write_records([row])
        result = self.inspect()
        export_inspection(result, self.output, detail_columns=(6, 8, 11, 12, 7, 9, 10, 14))
        wb = load_workbook(str(self.output))
        try:
            sheet = wb["項目明細"]
            headers = [cell.value for cell in sheet[1]]
            self.assertEqual(headers, ["FieldID", "項目名", "受領OCR値", "文字数",
                                       "最大桁数", "属性", "座標", "備考"])
            self.assertEqual(sheet["A2"].value, "6001")
            self.assertEqual(sheet["D2"].value, 21)
            self.assertEqual(sheet["E2"].value, 20)
            self.assertIn("最大桁数超過", sheet["H2"].value)
            self.assertEqual(sheet["H2"].fill.fgColor.rgb[-6:], "FFF2CC")
        finally:
            wb.close()
        with self.assertRaises(LayoutTxtError):
            export_inspection(result, self.output, overwrite=True, detail_columns=(8, 11, 12))

    def test_layout_definition_columns_can_be_selected_and_note_stays_last(self):
        wb = load_workbook(str(self.definition))
        try:
            sheet = wb.active
            for column, heading, value in (
                    (13, "入力属性", "必須"), (14, "入力規則", "英数字のみ"),
                    (15, "補足", "照合用"), (16, "出力例", "ABC123")):
                sheet.cell(1, column, heading)
                sheet.cell(2, column, value)
            wb.save(str(self.definition))
        finally:
            wb.close()
        self.write_records([self.row()])
        result = self.inspect()
        export_inspection(result, self.output,
                          detail_columns=(6, 8, 11, 12, 20, 21, 22, 23, 14))
        exported = load_workbook(str(self.output))
        try:
            sheet = exported["項目明細"]
            self.assertEqual([cell.value for cell in sheet[1]],
                             ["FieldID", "受領OCR値", "属性", "座標", "データ型", "IME",
                              "入力属性", "入力規則", "備考"])
            self.assertEqual([sheet.cell(2, col).value for col in range(5, 9)],
                             ["文字列", "半角英数", "必須", "英数字のみ"])
        finally:
            exported.close()
        with self.assertRaises(LayoutTxtError):
            export_inspection(result, self.output, overwrite=True,
                              detail_columns=(6, 8, 11, 12, 24))

    def test_generated_multiple_forms_including_calendar_round_trip(self):
        generated = generate_layout_txt(self.definition, self.root / "generated",
                                        error_patterns="none", generate_tif=False)
        result = inspect_txt(self.definition, generated.txt_files)
        self.assertEqual(result.exit_code, 0)
        self.assertEqual(len(result.records), 2)
        calendar = [row for row in result.details if row[3] == "2001"]
        self.assertEqual(len(calendar), 46)
        self.assertEqual(calendar[-1][6], "46")
        self.assertEqual(calendar[-1][19], "46/46")
        self.assertEqual(calendar[-1][18], "0200")

    def test_windows_xml_newline_translation_preserves_long_cell_and_marker_text(self):
        self.write_records([self.row()])
        result = self.inspect()
        value = "\U000f0000" + "\rA\nB\r\n" * 4000
        result.details[0][8] = value
        original_save = Workbook.save

        def windows_save(workbook, stream):
            original_save(workbook, stream)
            stream.seek(0)
            translated = io.BytesIO()
            with zipfile.ZipFile(stream) as source, zipfile.ZipFile(translated, "w") as target:
                for entry in source.infolist():
                    data = source.read(entry.filename)
                    if entry.filename.endswith(".xml"):
                        data = data.replace(b"\n", b"\r\n")
                    target.writestr(entry, data)
            stream.seek(0)
            stream.truncate()
            stream.write(translated.getvalue())

        with mock.patch.object(Workbook, "save", windows_save):
            export_inspection(result, self.output)
        wb = load_workbook(str(self.output))
        try:
            self.assertEqual(wb["項目明細"]["G2"].value, value)
            self.assertEqual(result.details[0][8], value)
        finally:
            wb.close()

    def test_extra_incomplete_block_is_retained_without_id_matching(self):
        self.write_records([["1001", "1", "1", "A", "0", "0,0,1,1",
                             "1", "B", "0", "0,0,1,1", "99", "C"]])
        result = self.inspect()
        self.assertEqual(result.exit_code, 0)
        self.assertEqual(len(result.details), 3)
        self.assertEqual([row[7] for row in result.details], ["受付番号", "氏名", ""])
        self.assertEqual(result.details[0][14], "")
        self.assertEqual(result.details[1][14], "")
        self.assertIn("余剰項目", result.details[2][14])
        self.assertIn("不足", result.details[2][14])
        self.assertEqual(result.details[2][8], "C")
        self.assertIsNone(result.details[2][11])

    def test_unknown_form_target_and_attributes_have_notes(self):
        rows = [self.row(), self.row(), self.row()]
        rows[0][0] = "9999"
        rows[1][1] = "9"
        rows[2][4] = "9"
        self.write_records(rows)
        self.assertEqual([r["status"] for r in self.inspect().records], ["注意点あり"] * 3)

    def test_received_attribute_values_0_4_8_12_are_valid(self):
        rows = []
        for attribute in ("0", "4", "8", "12"):
            row = self.row()
            row[4] = row[8] = attribute
            rows.append(row)
        self.write_records(rows)
        result = self.inspect()
        self.assertEqual([record["actual"] for record in result.records], [2] * 4)
        self.assertEqual([row[11] for row in result.details[::2]], ["0", "4", "8", "12"])
        self.assertEqual([row[14] for row in result.details], [""] * 8)
        self.assertEqual([record["issues"] for record in result.records], [""] * 4)
        export_inspection(result, self.output)
        wb = load_workbook(str(self.output))
        try:
            self.assertEqual([wb["項目明細"]["H%d" % row].value for row in range(2, 10)],
                             ["0", "0", "4", "4", "8", "8", "12", "12"])
            self.assertEqual(wb["項目明細"].max_column, 10)
        finally:
            wb.close()

    def test_received_attribute_12_without_coordinates_keeps_three_field_format(self):
        row = self.row()
        row[4] = row[8] = "12"
        self.write_records([row[:5] + row[6:9]])
        result = self.inspect()
        self.assertEqual(result.records[0]["actual"], 2)
        self.assertEqual([detail[8] for detail in result.details], ["0000123", "山田 太郎"])
        self.assertEqual([detail[11] for detail in result.details], ["12", "12"])
        self.assertEqual([detail[12] for detail in result.details], [None, None])
        self.assertEqual(result.records[0]["issues"], "")

    def test_old_attribute_values_are_not_accepted_as_received_format(self):
        row = self.row()
        row[4], row[8] = "1", "2"
        self.write_records([row])
        result = self.inspect()
        self.assertEqual(result.records[0]["actual"], 2)
        self.assertTrue(all("属性が定義外" in detail[14] for detail in result.details))

    def test_empty_values_and_invalid_attributes_require_review(self):
        row = self.row(first="")
        row[4] = "1,2"
        row[1] = "0"
        self.write_records([row])
        result = self.inspect()
        self.assertEqual(result.exit_code, 0)
        self.assertEqual(result.details[0][8], "")
        self.assertIn("属性が定義外", result.details[0][14])
        self.assertNotIn("原票", result.details[0][14])
        self.assertEqual(result.records[0]["status"], "注意点あり")

    def test_overlong_and_bad_coordinates(self):
        row = self.row(first="あ" * 21)
        row[-1] = "bad"
        self.write_records([row])
        result = self.inspect()
        self.assertEqual(result.exit_code, 0)
        self.assertEqual(result.details[0][8], "あ" * 21)
        self.assertIn("最大桁数超過", result.details[0][14])
        self.assertIn("座標", result.details[1][14])

    def test_blank_and_omitted_final_coordinate_have_no_notes(self):
        row = self.row()
        row[5] = ""
        self.write_records([row[:-1]])
        result = self.inspect()
        self.assertEqual(result.records[0]["issues"], "")
        self.assertEqual([r[14] for r in result.details], ["", ""])
        self.assertEqual([r[12] for r in result.details], ["", None])

    def test_no_coordinate_columns_preserve_all_field_values(self):
        row = self.row(first="000123")
        self.write_records([row[:5] + row[6:9]])
        result = self.inspect()
        self.assertEqual(result.records[0]["issues"], "")
        self.assertEqual(result.records[0]["actual"], 2)
        self.assertEqual([r[8] for r in result.details], ["000123", "山田 太郎"])
        self.assertEqual([r[12] for r in result.details], [None, None])
        export_inspection(result, self.output)
        wb = load_workbook(str(self.output))
        try:
            self.assertIsNone(wb["項目明細"]["I2"].value)
            self.assertEqual(wb["項目明細"]["G2"].value, "000123")
            self.assertFalse(any("NG" in str(cell.value) for sheet in wb for row in sheet for cell in row))
        finally:
            wb.close()

    def test_coordinate_formats_are_detected_per_record(self):
        row = self.row()
        self.write_records([row, row[:5] + row[6:9]])
        result = self.inspect()
        self.assertEqual([r["issues"] for r in result.records], ["", ""])
        self.assertEqual([r[8] for r in result.details], ["0000123", "山田 太郎"] * 2)
        self.assertEqual([r[12] for r in result.details], ["0,0,10,10", "0,0,10,10", None, None])

    def test_ambiguous_blocks_keep_raw_without_guessing_values(self):
        self.write_records([["9999", "1"] + ["0"] * 12])
        result = self.inspect()
        self.assertEqual(result.details, [])
        self.assertIsNone(result.records[0]["actual"])
        self.assertIn("区切りを特定できません", result.records[0]["issues"])
        self.assertEqual(result.records[0]["raw"], self.txt.read_bytes().decode("cp932"))

    def test_blank_lines_empty_file_and_malformed_csv_are_retained_with_notes(self):
        for payload in (b"", b"\r\n", b'"1001","1","unterminated\r\ntrailing data'):
            with self.subTest(payload=payload):
                self.txt.write_bytes(payload)
                result = self.inspect()
                self.assertEqual(result.exit_code, 0)
                self.assertEqual("".join(r["raw"] for r in result.records), payload.decode("cp932"))

    def test_utf8_bom_and_selected_encoding(self):
        self.write_records([self.row()], encoding="utf-8-sig")
        result = self.inspect(encoding="utf-8-sig")
        self.assertEqual(result.exit_code, 0)
        self.assertTrue(result.records[0]["raw"].startswith("\ufeff"))
        self.assertEqual(result.files[0][2], hashlib.sha256(self.txt.read_bytes()).hexdigest())
        self.txt.write_bytes(b"\x81")
        with self.assertRaises(LayoutTxtError):
            self.inspect()

    def test_input_sources_and_existing_output_are_protected(self):
        self.write_records([self.row()])
        result = self.inspect()
        before = self.definition.read_bytes()
        with self.assertRaises(LayoutTxtError):
            export_inspection(result, self.definition, overwrite=True)
        self.assertEqual(self.definition.read_bytes(), before)
        export_inspection(result, self.output)
        previous = self.output.read_bytes()
        with self.assertRaises(LayoutTxtError):
            export_inspection(result, self.output)
        self.assertEqual(self.output.read_bytes(), previous)
        export_inspection(result, self.output, overwrite=True)

    def test_invalid_xml_is_visible_and_excel_never_silently_truncates(self):
        self.write_records([self.row(first="A\x01B")])
        result = self.inspect()
        export_inspection(result, self.output)
        wb = load_workbook(str(self.output))
        try:
            self.assertEqual(wb["項目明細"]["G2"].value, "A\\u0001B")
        finally:
            wb.close()
        result.details[0][8] = "a" * 32768
        original = self.output.read_bytes()
        with self.assertRaises(LayoutTxtError):
            export_inspection(result, self.output, overwrite=True)
        self.assertEqual(self.output.read_bytes(), original)

    def test_nul_is_retained_in_raw_even_when_legacy_csv_rejects_it(self):
        self.write_records([self.row(first="A_NUL_B")])
        self.txt.write_bytes(self.txt.read_bytes().replace(b"_NUL_", b"\x00"))
        result = self.inspect()
        export_inspection(result, self.output)
        wb = load_workbook(str(self.output))
        try:
            self.assertIn("A\\u0000B", wb["受領原文"]["E2"].value)
            if sys.version_info < (3, 11):
                self.assertIn("CSV解析不可", result.records[0]["issues"])
            else:
                self.assertEqual(wb["項目明細"]["G2"].value, "A\\u0000B")
        finally:
            wb.close()

    def test_raw_long_records_are_split_without_loss(self):
        self.write_records([self.row(first="a" * 12000, second="b" * 12000)])
        result = self.inspect()
        export_inspection(result, self.output)
        wb = load_workbook(str(self.output))
        try:
            self.assertEqual("".join(row[0].value for row in wb["受領原文"].iter_rows(
                min_row=2, min_col=5, max_col=5)), self.txt.read_bytes().decode("cp932"))
        finally:
            wb.close()

    def test_duplicate_file_selection_and_empty_input_are_rejected(self):
        self.write_records([self.row()])
        for paths in ([], [self.txt, self.txt]):
            with self.assertRaises(LayoutTxtError):
                inspect_txt(self.definition, paths)

    def test_cli_exports_notes_without_failure_exit_code(self):
        args = [str(self.definition), str(self.txt), "--out", str(self.output)]
        for value, expected in (("A", 0), ("a" * 21, 0), ("", 0)):
            self.write_records([self.row(first=value)])
            self.assertEqual(main(args + ["--overwrite"]), expected)
            self.assertTrue(self.output.is_file())
        self.assertEqual(main(args), 2)
        row = self.row()
        self.write_records([row[:5] + row[6:9]])
        self.assertEqual(main(args + ["--overwrite", "--no-coordinates"]), 0)

    def test_native_window_parse_filter_raw_and_export(self):
        try:
            import tkinter as tk
        except ImportError:
            self.skipTest("tkinterなし")
        if not hasattr(tk, "Tk"):
            self.skipTest("tkinterなし")
        from autotest.layout_inspect_gui import InspectionWindow
        from autotest import layout_inspect_gui
        try:
            window = InspectionWindow()
        except tk.TclError:
            self.skipTest("表示環境なし")
        root = window.window
        self.write_records([self.row(), self.row(first="a" * 21)])
        window.window.withdraw()

        def wait_for_worker():
            deadline = time.monotonic() + 5
            while window.busy and time.monotonic() < deadline:
                root.update()
                time.sleep(0.01)
            self.assertFalse(window.busy)
        try:
            with mock.patch.object(layout_inspect_gui.filedialog, "askopenfilename",
                                   return_value=str(self.definition)):
                window._choose_excel()
            self.assertEqual(window.sheet.get(), "定義")
            with mock.patch.object(layout_inspect_gui.filedialog, "askopenfilenames",
                                   return_value=[str(self.txt)]):
                window._choose()
            wait_for_worker()
            self.assertEqual(len(window.record_tree.get_children()), 2)
            self.assertEqual(len(window.detail_tree.get_children()), 4)
            window.detail_tree.selection_set("d1", "d0")
            window._copy_rows(window.detail_tree)
            copied = list(csv.reader(io.StringIO(root.clipboard_get()), delimiter="\t"))
            self.assertEqual([row[6] for row in copied], ["0000123", "山田 太郎"])
            root.deiconify()
            window.detail_tree.master.master.select(window.detail_tree.master)
            root.update()
            tree = window.detail_tree
            x, y, width, height = tree.bbox("d0", "#7")
            # 最初のクリックとドラッグだけでセル内選択でき、別窓は作らない。
            tree.event_generate("<Button-1>", x=x + 12, y=y + 4)
            root.update()
            cell_text = window._cell_text
            self.assertIsNotNone(cell_text)
            self.assertIs(cell_text.master, tree)
            self.assertIs(cell_text.winfo_toplevel(), root)
            self.assertEqual(cell_text.get("1.0", "end-1c"), "0000123")
            self.assertEqual(str(cell_text["state"]), "disabled")
            from tkinter import font as tkfont
            font = tkfont.Font(root=root, font=cell_text["font"])
            tree.event_generate("<B1-Motion>", x=x + 12 + font.measure("000"), y=y + 4)
            tree.event_generate("<ButtonRelease-1>", x=x + 12 + font.measure("000"), y=y + 4)
            root.update()
            # 字体/DPIごとのピクセル境界差は許容し、部分選択とコピー一致を確認。
            selected_text = cell_text.get("sel.first", "sel.last")
            self.assertTrue(selected_text)
            self.assertLess(len(selected_text), len("0000123"))
            self.assertIn(selected_text, "0000123")
            cell_text.event_generate("<<Copy>>")
            root.update()
            self.assertEqual(root.clipboard_get(), selected_text)
            cell_text.insert("1.0", "changed")
            self.assertEqual(tree.set("d0", "#7"), "0000123")
            self.assertEqual(cell_text.get("1.0", "end-1c"), "0000123")
            x2, y2, _width, _height = tree.bbox("d1", "#7")
            toggle_modifier = 0x0008 if sys.platform == "darwin" else 0x0004
            tree.event_generate("<Button-1>", x=x2 + 12, y=y2 + 4, state=toggle_modifier)
            root.update()
            self.assertIsNone(window._cell_text)
            self.assertEqual(set(tree.selection()), {"d0", "d1"})
            window._copy_rows(tree)
            copied = list(csv.reader(io.StringIO(root.clipboard_get()), delimiter="\t"))
            self.assertEqual([row[6] for row in copied], ["0000123", "山田 太郎"])
            window._select_cell_text(tree, "d0", "#7")
            window._scroll_table(tree, "x", "moveto", 0)
            self.assertIsNone(window._cell_text)
            root.withdraw()
            window.issues_only.set(True)
            window._render()
            self.assertEqual(len(window.record_tree.get_children()), 1)
            self.assertEqual(len(window.detail_tree.get_children()), 1)
            window.record_tree.selection_set("r1")
            root.update()
            self.assertEqual(window.raw_text.get("1.0", "end-1c"), window.result.records[1]["raw"])
            with mock.patch.object(layout_inspect_gui.filedialog, "asksaveasfilename",
                                   return_value=str(self.output)):
                window._export()
            wait_for_worker()
            self.assertEqual(window.saved_path, self.output.resolve())
            wb = load_workbook(str(self.output))
            try:
                self.assertEqual(wb["項目明細"].max_row, 5)
            finally:
                wb.close()
            window._invalidate()
            self.assertIsNone(window.result)
            self.assertEqual(len(window.detail_tree.get_children()), 0)
            self.assertEqual(str(window.export_button["state"]), "disabled")
        finally:
            root.destroy()


if __name__ == "__main__":
    unittest.main()
